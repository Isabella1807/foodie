# Laver src/data/frida.json ud fra Den Danske Fødevaredatabases regneark (Frida).
#
# Regnearket hentes fra DTU Data: https://doi.org/10.11583/DTU.32312844
# (filen FCDB_x.y_Dataset.xlsx). Data er udgivet under CC BY 4.0, så kilden
# skal nævnes, hvor tallene vises — det står i filens "credit".
#
# Brug:  python scripts/frida_to_json.py sti/til/FCDB_6.1_Dataset.xlsx
# Kræver: pip install openpyxl
#
# Vi tager "deklarations"-tallene (dem der må stå på en pakke), så de svarer
# til tallene fra Open Food Facts og fra etiketter: kcal, protein, tilgængeligt
# kulhydrat, fedt og kostfibre — alt pr. 100 gram.
import json
import re
import sys

import openpyxl

if len(sys.argv) < 2:
    sys.exit("Brug: python scripts/frida_to_json.py FCDB_6.1_Dataset.xlsx")

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True)

# Version og dato står i Readme-arket, fx "... version 6.1, Maj 2026"
version, date = "?", ""
for row in wb["Readme"].iter_rows(values_only=True):
    m = re.search(r"version\s+([\d.]+),\s*([A-Za-zæøåÆØÅ]+\s+\d{4})", str(row[0] or ""))
    if m:
        version, date = m.group(1), m.group(2).lower()
        break

# Fødevaregruppe pr. FoodID (fx "Bærfrugt"), til at skelne ens navne
groups = {}
food_rows = wb["Food"].iter_rows(values_only=True)
next(food_rows)  # overskrifter
for r in food_rows:
    if r[2] is not None:
        groups[str(r[2])] = r[11]

# Data_Table: første fire rækker er overskrifter (navn, engelsk navn, enhed, id)
COL = {"name": 0, "id": 2, "kcal": 6, "protein": 9, "carbs": 12, "fiber": 13, "fat": 14}
rows = wb["Data_Table"].iter_rows(values_only=True)
for _ in range(4):
    next(rows)


def num(v, digits):
    if v is None or v == "":
        return None
    try:
        return round(float(v), digits) if digits else round(float(v))
    except ValueError:
        return None


foods = []
for r in rows:
    name = r[COL["name"]]
    kcal = num(r[COL["kcal"]], 0)
    if not name or kcal is None:
        continue
    foods.append([
        str(name).strip(),
        kcal,
        num(r[COL["protein"]], 1),
        num(r[COL["carbs"]], 1),
        num(r[COL["fat"]], 1),
        num(r[COL["fiber"]], 1),
        groups.get(str(r[COL["id"]])) or None,
    ])

foods.sort(key=lambda f: f[0].lower())

out = {
    "name": "Fødevaredatabasen (DTU)",
    "credit": f"Den Danske Fødevaredatabase (fcdb.fooddata.dk), version {version}, {date}, DTU Fødevareinstituttet. CC BY 4.0.",
    "version": version,
    "fields": ["name", "kcal", "protein", "carbs", "fat", "fiber", "group"],
    "per": "100 g",
    "foods": foods,
}

with open("src/data/frida.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

print(f"{len(foods)} fødevarer skrevet til src/data/frida.json (version {version}, {date})")
